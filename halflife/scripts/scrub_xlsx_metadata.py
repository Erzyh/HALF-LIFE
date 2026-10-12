"""Remove person-identifying metadata from .xlsx workbooks without touching any cell.

Applied in place on 2026-10-11 to every returned workbook under coding/results*/ (the coders and judges are anonymised A to D
in the paper, but Excel had recorded their account names and the folder they saved from). Cell content was unchanged.

    python scrub_xlsx_metadata.py --out DIR file.xlsx ...   write scrubbed copies into DIR (originals untouched) and verify
    python scrub_xlsx_metadata.py --in-place file.xlsx ...  rewrite the files themselves and verify against the original bytes

What is removed
  docProps/core.xml   <dc:creator> and <cp:lastModifiedBy> are emptied
  xl/workbook.xml     the <mc:AlternateContent> block that carries <x15ac:absPath url="C:\\Users\\<name>\\..."/> is dropped
                      (Excel writes the folder the file was last saved in; it is optional and Excel re-creates it on save)
Everything else is copied byte for byte. Verification loads the original and the scrubbed workbook with openpyxl and
compares sheet names, dimensions and every cell value; the script exits non-zero if anything differs.
"""
import argparse
import io
import re
import shutil
import sys
import zipfile
from pathlib import Path

import openpyxl

CORE_SUBS = [
    (re.compile(r"<dc:creator>.*?</dc:creator>", re.S), "<dc:creator></dc:creator>"),
    (re.compile(r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>", re.S), "<cp:lastModifiedBy></cp:lastModifiedBy>"),
]
ALT_BLOCK = re.compile(r"<mc:AlternateContent\b.*?</mc:AlternateContent>", re.S)


def scrub_bytes(data: bytes):
    """Return (scrubbed zip bytes, list of changes)."""
    changes = []
    src = zipfile.ZipFile(io.BytesIO(data))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as dst:
        for info in src.infolist():
            raw = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                txt = raw.decode("utf-8")
                for pat, rep in CORE_SUBS:
                    txt, n = pat.subn(rep, txt)
                    if n and rep not in raw.decode("utf-8"):
                        changes.append(pat.pattern.split(">")[0].lstrip("<"))
                raw = txt.encode("utf-8")
            elif info.filename == "xl/workbook.xml":
                txt = raw.decode("utf-8")
                def drop(m):
                    if "absPath" in m.group(0):
                        changes.append("absPath")
                        return ""
                    return m.group(0)
                raw = ALT_BLOCK.sub(drop, txt).encode("utf-8")
            new = zipfile.ZipInfo(info.filename, date_time=info.date_time)
            new.compress_type = info.compress_type
            new.external_attr = info.external_attr
            dst.writestr(new, raw)
    return buf.getvalue(), changes


def snapshot(data: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(data))
    out = []
    for ws in wb.worksheets:
        cells = tuple((c.coordinate, c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        out.append((ws.title, ws.max_row, ws.max_column, cells))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+")
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--out", help="directory for scrubbed copies")
    g.add_argument("--in-place", action="store_true")
    a = ap.parse_args()
    bad = 0
    for f in a.files:
        p = Path(f)
        data = p.read_bytes()
        scrubbed, changes = scrub_bytes(data)
        same = snapshot(data) == snapshot(scrubbed)
        if not same:
            bad += 1
        if a.in_place:
            if same:
                p.write_bytes(scrubbed)
        else:
            dst = Path(a.out) / p.name
            dst.parent.mkdir(parents=True, exist_ok=True)
            dst.write_bytes(scrubbed)
        print(f"{p.as_posix()}: removed {sorted(set(changes)) or 'nothing'}; cells identical: {same}; {len(data):,} -> {len(scrubbed):,} bytes")
    if bad:
        print(f"{bad} file(s) differ after scrubbing; those were NOT written in place", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
