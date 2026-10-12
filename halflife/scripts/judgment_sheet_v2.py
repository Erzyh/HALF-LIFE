"""Judgment workbooks, version 2 (6 October 2026): same rows and condition codes as version 1, but every mechanical
step done in advance so that a judge answers two plain questions per row.

For each row the sheet shows, in Korean, what was measured, the paper's value, the re-run value with its interval,
whether the paper's value lies inside the interval, the difference in points on the paper's 0-100 scale (fractions
converted; non-score metrics labelled), and one sentence stating what the paper said with that number. The judge
fills two dropdowns (판정, 결론 수준) and an optional memo. The automatic class is NOT shown (independence); the
merge script takes it from the version-1 master. Row order and the A-E condition codes are those of
coding/judgment_sheet_v1.xlsx, verified against coding/judgment_key_v1.csv, so the key stays valid.

python scripts/judgment_sheet_v2.py            -> coding/judgment_v2/halflife_judgment_v2_judge{A,B,C}.xlsx
                                                  and coding/judgment_sheet_v2.xlsx (master: all three judges + automatic class)
Inputs: coding/judgment_sheet_v1.xlsx, coding/judgment_key_v1.csv, coding/judgment_sentences_v2.csv,
        harness/tierb/configs/<Wid>.json (model id). No API calls.
"""
import csv
import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HALF = Path(__file__).resolve().parents[1]
V1 = HALF / "coding" / "judgment_sheet_v1.xlsx"
KEY = HALF / "coding" / "judgment_key_v1.csv"
SENT = HALF / "coding" / "judgment_sentences_v2.csv"
OUT_DIR = HALF / "coding" / "judgment_v2"
MASTER = HALF / "coding" / "judgment_sheet_v2.xlsx"
FONT = "Malgun Gothic"
YELLOW = PatternFill("solid", fgColor="FFFF00")
GREY = PatternFill("solid", fgColor="EEEEEE")
JUDG = '"생존,부분 생존,사망,판정 불능"'
CONCL = '"핵심 주장 성립,부분 성립,불성립,판단 불가"'
VERSION = "v2, 2026-10-06"
MODEL_NAME = {"gpt-4-turbo-2024-04-09": "GPT-4 Turbo", "gpt-4o-2024-05-13": "GPT-4o", "o1-2024-12-17": "o1", "o3-mini-2025-01-31": "o3-mini",
              "o4-mini-2025-04-16": "o4-mini", "gpt-4.1-nano-2025-04-14": "GPT-4.1 nano", "gpt-5-2025-08-07": "GPT-5", "gpt-5-mini-2025-08-07": "GPT-5 mini",
              "o3-2025-04-16": "o3", "o1-pro-2025-03-19": "o1-pro", "gpt-5-nano-2025-08-07": "GPT-5 nano", "gpt-4.1-nano": "GPT-4.1 nano"}
BATCHIM = {"o1", "모델"}  # names whose Korean reading ends in a consonant (o1 = 오원)


def fill(sentence, name, value):
    b = name in BATCHIM
    return (sentence.replace("{모델}은", name + ("은" if b else "는")).replace("{모델}이 ", name + ("이 " if b else "가 "))
            .replace("{모델}", name).replace("{값}", value))


def fmt(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    s = f"{float(v):.4g}" if abs(float(v)) < 1 else f"{float(v):.6g}"
    return s


def model_of(wid):
    p = HALF / "harness" / "tierb" / "configs" / f"{wid}.json"
    if p.exists():
        mid = (json.load(open(p, encoding="utf-8")).get("model") or {}).get("id", "")
        return MODEL_NAME.get(mid, mid or "모델")
    return "모델"


def pts_text(x):
    """Points with one decimal, two when one decimal would move the value across the 2- or 5-point threshold."""
    r1 = round(x, 1)
    if any((x <= th) != (r1 <= th) for th in (2.0, 5.0)):
        return f"{x:.2f}점"
    return f"{r1:.1f}점"


def diff_prompt(o, r, lo, hi, nonscore):
    """(inside, diff_text) for a prompt-level row."""
    if r is None or pd.isna(r):
        return "값 없음", "값 없음"
    o, r = float(o), float(r)
    inside = "없음" if lo is None or pd.isna(lo) else ("안" if float(lo) <= o <= float(hi) else "밖")
    if nonscore:
        return inside, f"{abs(r - o):.3g} (점수 아님)"
    unit = 100.0 if abs(o) <= 1 and abs(r) <= 1 else 1.0
    return inside, pts_text(abs(r - o) * unit)


def diff_tiera(o, r, lo, hi, rule):
    if o is None or r is None or pd.isna(o) or pd.isna(r):
        return "값 없음", "값 없음"
    o, r = float(o), float(r)
    inside = "없음" if lo is None or pd.isna(lo) else ("안" if float(lo) <= o <= float(hi) else "밖")
    rule = str(rule)
    if rule == "x100":
        return inside, pts_text(abs(r - o) * 100)
    if rule == "pts":
        return inside, pts_text(abs(r - o))
    return inside, f"{abs(r - o):.3g} (점수 아님: {rule.replace('other:', '')})"


def build(judges, with_auto):
    v1a = pd.read_excel(V1, sheet_name="TierA_claims")
    v1b = pd.read_excel(V1, sheet_name="PromptLevel_blind")
    key = list(csv.DictReader(open(KEY, encoding="utf-8-sig")))
    assert [(int(k["row"]), k["wid"], k["code"]) for k in key] == [(int(r["번호"]), r["논문 ID"], r["조건 부호"]) for _, r in v1b.iterrows()], "v1 master and key differ"
    sent = {r["id"]: r for r in csv.DictReader(open(SENT, encoding="utf-8-sig"))}
    wb = Workbook()
    # ---- PromptLevel_blind
    ws = wb.active
    ws.title = "PromptLevel_blind"
    head = ["번호", "논문 ID", "논문", "지표", "원 논문 값", "재실행 값", "95% 구간 하한", "95% 구간 상한", "오차 범위 안?", "차이", "비고", "조건 부호"]
    if with_auto:
        head.append("자동 가안(값 수준)")
    for j in judges:
        head += [f"판정 {j}"]
    head.append("논문이 한 말")
    for j in judges:
        head += [f"결론 수준 {j}", f"메모 {j}"]
    ws.append(head)
    models = {}
    for _, r in v1b.iterrows():
        wid = r["논문 ID"]
        s = sent.get(wid, {})
        models.setdefault(wid, model_of(wid))
        nonscore = s.get("방향", "") == "점수 아님"
        inside, dtxt = diff_prompt(r["원 논문 값"], r["재실행 값"], r["95% 구간 하한"], r["95% 구간 상한"], nonscore)
        auto = str(r["자동 가안(값 수준)"])
        if auto == "확인 필요":
            note = f"응답 {int(r['문항 수'])}개 중 {int(r['파싱 실패 수'])}개 파싱 안 됨 (절반 넘음): 판정 불능"
        elif auto == "해당없음":
            note = "과제 특성상 이 행에는 적용 불가: 판정 불능"
        else:
            note = ""
        what = s.get("무엇을 쟀나", r["지표"])
        said = fill(s.get("논문이 한 말", ""), models[wid], fmt(r["원 논문 값"]))
        row = [int(r["번호"]), wid, r["논문"], what, r["원 논문 값"], r["재실행 값"], r["95% 구간 하한"], r["95% 구간 상한"], inside, dtxt, note, r["조건 부호"]]
        if with_auto:
            row.append(r["자동 가안(값 수준)"])
        row += [None] * len(judges) + [said] + [None, None] * len(judges)
        ws.append(row)
    n_b = len(v1b)
    style(ws, head, judges, n_b, with_auto, prompt=True)
    # ---- TierA_claims
    wt = wb.create_sheet("TierA_claims")
    head_a = ["번호", "카드", "논문", "주장(지표)", "원 논문 값", "재실행 값", "95% 구간 하한", "95% 구간 상한", "오차 범위 안?", "차이", "사람 확인 항목"]
    if with_auto:
        head_a.append("자동 판정(논문 단위 가안)")
    for j in judges:
        head_a += [f"판정 {j}"]
    head_a.append("논문이 한 말")
    for j in judges:
        head_a += [f"결론 수준 {j}", f"메모 {j}"]
    wt.append(head_a)
    for _, r in v1a.iterrows():
        s = sent.get(r["카드"], {})
        inside, dtxt = diff_tiera(r["원 논문 값"], r["재실행 값"], r["95% 구간 하한"], r["95% 구간 상한"], r["규칙(척도)"])
        said = s.get("논문이 한 말", "")
        said = (said + f" / 이 행: {r['주장(지표)']}") if said else f"이 행: {r['주장(지표)']}"
        row = [int(r["번호"]), r["카드"], r["논문"], r["주장(지표)"], r["원 논문 값"], r["재실행 값"], r["95% 구간 하한"], r["95% 구간 상한"], inside, dtxt,
               r["사람 확인 항목"] if isinstance(r["사람 확인 항목"], str) else ""]
        if with_auto:
            row.append(r["자동 판정(논문 단위 가안)"])
        row += [None] * len(judges) + [said] + [None, None] * len(judges)
        wt.append(row)
    n_a = len(v1a)
    style(wt, head_a, judges, n_a, with_auto, prompt=False)
    guide(wb, n_a, n_b, judges[0] if len(judges) == 1 else None)
    return wb, n_a, n_b


def style(ws, head, judges, nrows, with_auto, prompt):
    ncol = len(head)
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, bold=True)
        cell.fill = GREY
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "E2"
    widths = {"번호": 6, "논문 ID": 13, "카드": 13, "논문": 34, "지표": 40, "주장(지표)": 40, "원 논문 값": 10, "재실행 값": 10, "95% 구간 하한": 10, "95% 구간 상한": 10,
              "오차 범위 안?": 9, "차이": 14, "비고": 22, "조건 부호": 7, "사람 확인 항목": 12, "자동 가안(값 수준)": 12, "자동 판정(논문 단위 가안)": 18, "논문이 한 말": 70}
    dv1 = DataValidation(type="list", formula1=JUDG, allow_blank=True)
    dv2 = DataValidation(type="list", formula1=CONCL, allow_blank=True)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    for i, h in enumerate(head, 1):
        L = get_column_letter(i)
        ws.column_dimensions[L].width = widths.get(h, 14 if h.startswith(("판정", "결론")) else 30 if h.startswith("메모") else 12)
        if h.startswith("판정 "):
            dv1.add(f"{L}2:{L}{nrows + 1}")
        if h.startswith("결론 수준 "):
            dv2.add(f"{L}2:{L}{nrows + 1}")
    for r in range(2, nrows + 2):
        for i, h in enumerate(head, 1):
            cell = ws.cell(row=r, column=i)
            cell.font = Font(name=FONT)
            cell.alignment = Alignment(wrap_text=h in ("논문", "지표", "주장(지표)", "논문이 한 말", "비고") or h.startswith("메모"), vertical="top")
            if h.startswith(("판정 ", "결론 수준 ", "메모 ")):
                cell.fill = YELLOW


def guide(wb, n_a, n_b, judge):
    ws = wb.create_sheet("안내", 0)
    who = f"이 파일은 판정자 {judge}용입니다. " if judge else ""
    lines = [
        f"재실행 결과 3인 판정 시트 ({VERSION})" + (f" - 판정자 {judge}" if judge else ""),
        "",
        who + "행마다 노란 칸 둘(판정, 결론 수준)만 고르면 됩니다. 메모는 하고 싶은 말이 있을 때만 적습니다. 계산은 하지 않습니다. 시트가 이미 다 해 두었습니다.",
        "",
        "질문 1 (판정): 논문의 값이 다시 나왔습니까? 보는 열은 '오차 범위 안?', '차이', '비고' 셋입니다.",
        "  오차 범위 안 → 생존.  밖이면 '차이'로: 2점 이내 → 생존, 5점 이내 → 부분 생존, 5점 초과 → 사망.",
        "  오차 범위가 '없음'인 행(구간이 계산되지 않은 행)은 판정 불능이 아닙니다. '차이'만으로 정합니다: 2점 이내 → 생존, 5점 이내 → 부분 생존, 5점 초과 → 사망.",
        "  '차이'가 '값 없음'이거나 '비고'에 '판정 불능'이라고 적혀 있으면 → 판정 불능.",
        "  '차이'에 (점수 아님)이 붙은 행만 점수 기준을 쓰지 않습니다: 오차 범위 안 → 생존, 밖 → 사망. 이 행들에 한해 오차 범위가 없으면 → 판정 불능.",
        "  이 기준과 다르게 보시면 그렇게 판정하고 메모에 이유를 적어 주세요. 이유가 적힌 행만 기준과 달라도 됩니다.",
        "",
        "질문 2 (결론 수준): '논문이 한 말' 열의 문장이 다시 잰 값으로도 맞습니까?",
        "  그대로 맞으면 핵심 주장 성립, 단서를 달아야 맞으면 부분 성립, 틀리면 불성립, 이 행의 값만으로 알 수 없으면 판단 불가. 판정이 판정 불능이면 판단 불가.",
        "  문장 끝의 괄호가 안내입니다. (문제 지적)·(편향 지적)·(…라고 지적): 다시 잰 값에서 그 문제가 그대로이거나 더 크면 성립, 문제가 사라졌으면 불성립.",
        "  (수치 보고)·(행동 보고): 숫자 자체가 주장이므로 판정을 따릅니다. 생존 → 핵심 주장 성립, 부분 생존 → 부분 성립, 사망 → 불성립.",
        "  괄호가 없는 문장은 성능 보고입니다. 다시 잰 값이 비슷하면 성립, 조금 다르면 부분 성립, 많이 다르면 불성립.",
        "",
        f"PromptLevel_blind: 프롬프트 수준 결과 {n_b}행(86편). 같은 논문의 행들은 사망 전 고정 모델 실행과 후속 모델 실행이 섞여 있고 조건은 A~E 부호로 가려져 있습니다. 조건을 추측하지 말고 행마다 값만 보고 답합니다. 같은 논문의 행은 '논문이 한 말'이 같지만 다시 잰 값이 다르므로 행마다 따로 답합니다.",
        f"TierA_claims: 저장소 기반 재실행 논문의 주장 {n_a}건. '논문이 한 말'은 논문 전체의 문장이고 '/ 이 행:' 뒤가 이 행의 주장입니다. 이 탭에서는 '차이'가 5점을 넘어도 방향이 같으면 부분 생존입니다. 사망은 방향이 뒤집히거나(부호가 바뀌거나) '논문이 한 말'이 더는 맞지 않을 때만 씁니다. '사람 확인 항목'이 '예'인 행은 판단이 어려우면 판정 불능으로 두어도 됩니다.",
        "",
        "다른 판정자와 상의하지 않습니다. 파일명을 바꾸지 않습니다. 열쇠(coding/judgment_key_v1.csv)는 판정이 끝날 때까지 열지 않습니다.",
        "",
        "예시 (형식만 보이기 위한 가상의 값):",
    ]
    for l in lines:
        ws.append([l])
    ex = judge or "A"
    ws.append(["번호", "지표", "원 논문 값", "재실행 값", "오차 범위 안?", "차이", f"판정 {ex}", "논문이 한 말", f"결론 수준 {ex}", f"메모 {ex}"])
    ws.append([0, "MMLU-Pro 정확도 (%)", 63.7, 61.2, "밖", "2.5점", "부분 생존", "GPT-4 Turbo는 MMLU-Pro에서 정확도 63.7%를 낸다", "부분 성립", ""])
    ws.append([0, "과잉 거부율 (%)", 12.7, 26.5, "밖", "13.8점", "사망", "GPT-4 Turbo는 무해한 요청의 12.7%를 거부한다 (문제 지적)", "핵심 주장 성립", "문제가 더 커짐"])
    ws.append([0, "계획 정확도 (%)", 5.0, 62.0, "밖", "57.0점", "사망", "GPT-4 Turbo는 유효한 계획을 5%만 만든다 (계획을 못 한다고 지적)", "불성립", "이제는 계획을 함"])
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name=FONT, bold=(cell.row in (1, len(lines) + 1)))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for rr in (len(lines) + 2, len(lines) + 3, len(lines) + 4):
        for c in ("G", "I", "J"):
            ws[f"{c}{rr}"].fill = YELLOW
    ws.column_dimensions["A"].width = 110
    for c in "BCDEFGHIJ":
        ws.column_dimensions[c].width = 18


def main():
    OUT_DIR.mkdir(exist_ok=True)
    wb, n_a, n_b = build(("A", "B", "C"), with_auto=True)
    wb.save(MASTER)
    print(f"master -> {MASTER.relative_to(HALF)} (TierA {n_a}, PromptLevel {n_b})")
    for j in ("A", "B", "C"):
        wb, _, _ = build((j,), with_auto=False)
        out = OUT_DIR / f"halflife_judgment_v2_judge{j}.xlsx"
        wb.save(out)
        print(f"judge {j} -> {out.relative_to(HALF)}")


if __name__ == "__main__":
    main()
