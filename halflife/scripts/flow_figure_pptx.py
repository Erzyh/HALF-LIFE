"""Extended Data Fig. 1 (flow of papers through the study) as an editable PowerPoint drawing.

Builds figures/flow_diagram.pptx from the same box texts as scripts/flow_figure.py (COLS): every box, arrow, side note and
column title is a separate shape, so the figure can be adjusted in PowerPoint and exported again. All boxes have the same
size (the height fits the longest wrapped text), their text is centred, the boxes of a column are spread evenly over
the slide height, every line stays inside the slide and no shape carries a shadow. With PowerPoint installed,
scripts/pptx_export.py renders the slide to figures/flow_diagram.png at 300 dpi. No API calls.
Usage: python scripts/flow_figure_pptx.py
"""
import math
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from flow_figure import COLS  # noqa: E402

INK, GREY, ACCENT, WHITE = RGBColor(0x11, 0x11, 0x11), RGBColor(0x6B, 0x6B, 0x6B), RGBColor(0xE8, 0xEE, 0xF4), RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"
W = 14.7                                          # slide width (in); the height follows the longest column
BOX_W, PITCH, X0 = 3.8, 4.8, 0.2                  # box width, column pitch, left margin (in)
HEAD_PT, DET_PT, NOTE_PT, TITLE_PT = 9, 8, 8, 11
CPL_HEAD, CPL_DET = 50, 64                        # characters per line at those sizes inside the box
TOP, BOTTOM, GAP = 0.95, 0.3, 0.75                # first box top, bottom margin, room between boxes for the arrow and a side note (in)


def wrapped_lines(text, cpl):
    return sum(max(1, math.ceil(len(line) / cpl)) for line in text.split("\n")) if text else 0


def box_height(head, detail):
    return 0.16 + 0.19 * wrapped_lines(head, CPL_HEAD) + (0.05 + 0.165 * wrapped_lines(detail, CPL_DET) if detail else 0.0)


def set_run(run, size, bold=False, color=INK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_text(tf, lines, size, bold=False, color=GREY, align=PP_ALIGN.LEFT, first_space=0):
    """Paragraphs into a text frame (the first paragraph already exists)."""
    for k, line in enumerate(lines):
        p = tf.paragraphs[0] if k == 0 and not tf.paragraphs[0].runs else tf.add_paragraph()
        if k == 0 and first_space:
            p.space_before = Pt(first_space)
        set_run(p.add_run(), size, bold, color)
        p.runs[-1].text = line
        p.alignment = align


def no_shadow(shape):
    """An explicit empty effect list: nothing inherits the theme's shadow."""
    spPr = shape._element.spPr
    if spPr.find(qn("a:effectLst")) is None:
        spPr.append(spPr.makeelement(qn("a:effectLst"), {}))


def main():
    prs = Presentation()
    cols = list(COLS.items())
    n_max = max(len(b) for _, b in cols)
    box_h = max(box_height(h, d) for _, boxes in cols for h, d, _, _ in boxes)      # one size for every box
    H = round(TOP + n_max * (box_h + GAP) + BOTTOM, 1)
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    def connector(x1, y1, x2, y2, color, width, arrow=False):
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = color
        c.line.width = Pt(width)
        if arrow:
            ln = c.line._get_or_add_ln()
            ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
        no_shadow(c)
        return c

    for k, (title, boxes) in enumerate(cols):
        x_left = X0 + k * PITCH
        xc = x_left + BOX_W / 2
        tb = slide.shapes.add_textbox(Inches(x_left), Inches(0.2), Inches(BOX_W), Inches(0.35))
        tb.text_frame.word_wrap = True
        add_text(tb.text_frame, [title], TITLE_PT, bold=True, color=INK, align=PP_ALIGN.CENTER)
        no_shadow(tb)
        connector(x_left, 0.62, x_left + BOX_W, 0.62, INK, 0.75)
        n = len(boxes)
        step = (H - TOP - BOTTOM - box_h) / max(n - 1, 1)             # boxes spread evenly from the top to the bottom of the column
        for i, (head, detail, note, accent) in enumerate(boxes):
            top = TOP + i * step
            shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x_left), Inches(top), Inches(BOX_W), Inches(box_h))
            shp.fill.solid()
            shp.fill.fore_color.rgb = ACCENT if accent else WHITE
            shp.line.color.rgb = INK
            shp.line.width = Pt(0.75)
            shp.shadow.inherit = False
            no_shadow(shp)
            tf = shp.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = tf.margin_right = Inches(0.08)
            tf.margin_top = tf.margin_bottom = Inches(0.05)
            add_text(tf, [head], HEAD_PT, bold=True, color=INK, align=PP_ALIGN.CENTER)
            if detail:
                add_text(tf, detail.split("\n"), DET_PT, color=INK, align=PP_ALIGN.CENTER, first_space=3)
            if i < n - 1:
                y_next = top + step
                connector(xc, top + box_h + 0.02, xc, y_next - 0.02, INK, 0.75, arrow=True)
                if note:
                    ym = (top + box_h + y_next) / 2
                    connector(xc, ym, xc + 0.2, ym, GREY, 0.6)
                    note_w = min(2.6, W - (xc + 0.28) - 0.15)
                    nb = slide.shapes.add_textbox(Inches(xc + 0.28), Inches(ym - GAP / 2), Inches(note_w), Inches(GAP))
                    nb.text_frame.word_wrap = True
                    nb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                    nb.text_frame.margin_left = nb.text_frame.margin_right = Inches(0.02)
                    nb.text_frame.margin_top = nb.text_frame.margin_bottom = Inches(0)
                    add_text(nb.text_frame, note.split("\n"), NOTE_PT, color=GREY)
                    no_shadow(nb)
    out = ROOT / "figures" / "flow_diagram.pptx"
    prs.save(out)
    print(f"-> {out.relative_to(ROOT)} ({W} x {H} in, boxes {BOX_W} x {box_h:.2f} in, {len(slide.shapes)} shapes)")


if __name__ == "__main__":
    main()
